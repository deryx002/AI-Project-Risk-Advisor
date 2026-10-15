
import csv
from pathlib import Path
from io import StringIO

import fitz  # PyMuPDF
import pandas as pd
from docx import Document


def extract_text(file_path: Path) -> str:
    """Extract text from a supported project document."""

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        with fitz.open(file_path) as pdf:
            text = "\n".join(page.get_text() for page in pdf)

    elif extension == ".docx":
        doc = Document(file_path)

        paragraphs = [p.text for p in doc.paragraphs]

        tables = []
        for table in doc.tables:
            for row in table.rows:
                tables.append(" | ".join(cell.text for cell in row.cells))

        text = "\n".join(paragraphs + tables)

    elif extension == ".csv":
        dataframe = pd.read_csv(file_path)
        text = dataframe.to_csv(index=False)

    elif extension == ".txt":
        text = file_path.read_text(encoding="utf-8-sig")

    else:
        raise ValueError(f"Unsupported file type: {extension}")

    if not text.strip():
        raise ValueError(
            "No text was extracted. The file may be empty, "
            "scanned, or unsupported."
        )

    return text.strip()
